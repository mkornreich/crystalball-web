"""Open an edited workbook with openpyxl (an independent reader) and print what it sees.

    python3 tools/check-edited-workbook.py <edited.xlsx>
"""

import sys

import openpyxl


def main():
    wb = openpyxl.load_workbook(sys.argv[1])
    ws = wb["Model"]
    print("C6", ws["C6"].value, "| D9", ws["D9"].value, "| B12", ws["B12"].value, "| C9", ws["C9"].value)
    print("col B width", ws.column_dimensions["B"].width, "| row 12 height", ws.row_dimensions[12].height)
    for dv in ws.data_validations.dataValidation:
        print("rule", dv.sqref, dv.type, dv.operator, dv.formula1, dv.errorStyle, dv.promptTitle, dv.showDropDown)
    print("calc", wb.calculation.fullCalcOnLoad)


main()
