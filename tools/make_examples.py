"""Generate clean example workbooks for the Crystal Ball Web add-in.

These reproduce the *public* model structure of two classic Oracle Crystal Ball
tutorials (Futura Apartments and Vision Research) so you can open them in Excel,
define the listed assumptions/forecasts in the add-in, and reproduce Crystal
Ball's published results. The files are generated from scratch (formulas only) —
they do not contain or redistribute Oracle's proprietary example files.

Usage: python3 tools/make_examples.py   ->  examples/*.xlsx
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

OUT = os.path.join(os.path.dirname(__file__), "..", "examples")
GREEN = PatternFill("solid", fgColor="C6E8C6")  # assumption
BLUE = PatternFill("solid", fgColor="BBD7F0")  # forecast
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)


def style_header(ws):
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 16


def setup_sheet(wb, lines):
    ws = wb.create_sheet("Setup (read me)")
    ws.column_dimensions["A"].width = 100
    for i, line in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=line)
        if line and not line.startswith(" ") and line.endswith(":"):
            c.font = BOLD
        c.alignment = Alignment(wrap_text=True, vertical="top")


def futura():
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    style_header(ws)
    ws["B1"] = "Futura Apartments"
    ws["B1"].font = TITLE
    ws["B5"] = "Number of Rental Units"
    ws["C5"] = 35
    ws["C5"].fill = GREEN
    ws["B6"] = "Rent per Unit"
    ws["C6"] = 500
    ws["B7"] = "Monthly Expenses"
    ws["C7"] = 15000
    ws["C7"].fill = GREEN
    ws["B9"] = "Profit or Loss"
    ws["B9"].font = BOLD
    ws["C9"] = "=C5*C6-C7"
    ws["C9"].fill = BLUE
    ws["C9"].font = BOLD

    setup_sheet(
        wb,
        [
            "Futura Apartments — how to run this in the Crystal Ball Web add-in:",
            "",
            "Assumptions (green cells) — select the cell, click “+ Assumption”, then set:",
            "   C5  Number of Rental Units   ->  Discrete Uniform, Minimum 30, Maximum 40",
            "   C7  Monthly Expenses          ->  Normal, Mean 15000, Std. Dev. 1000",
            "",
            "Forecast (blue cell):",
            "   C9  Profit or Loss            ->  click “+ Forecast”",
            "",
            "Run settings (Run tab):",
            "   Trials 1000, Monte Carlo, fixed seed 999",
            "",
            "Expected result (matches Oracle Crystal Ball):",
            "   Mean profit ≈ $2,500; certainty of Profit ≥ $0 ≈ 90% (Crystal Ball reports 90.09%).",
            "   Set the certainty field to 90% to read the profit range.",
        ],
    )
    wb.save(os.path.join(OUT, "futura_apartments.xlsx"))
    print("wrote examples/futura_apartments.xlsx")


def vision_research():
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    style_header(ws)
    ws["B1"] = "Vision Research — ClearView Project"
    ws["B1"].font = TITLE
    ws["B2"] = "(in $millions)"
    rows = [
        ("B3", "Costs", None, None),
        ("B4", "Development Cost to Date", "C4", 10),
        ("B5", "Testing Costs", "C5", 4),
        ("B6", "Marketing Costs", "C6", 16),
        ("B7", "Total Costs", "C7", "=SUM(C4:C6)"),
        ("B9", "Drug Test (sample of 100)", None, None),
        ("B10", "Patients Cured", "C10", 25),
        ("B11", "FDA Approved if 20+ Cured", "C11", "=C10>=20"),
        ("B13", "Market Study", None, None),
        ("B14", "Persons Nearsighted Today (MM)", "C14", 40),
        ("B15", "Growth Rate of Nearsightedness", "C15", 0.02),
        ("B16", "Persons Nearsighted After 1 Year", "C16", "=C14*(1+C15)"),
        ("B19", "Market Penetration", "C19", 0.08),
        ("B20", "Profit per Customer ($)", "C20", 12),
        ("B21", "Gross Profit if Approved (MM)", "C21", "=C16*C19*C20"),
        ("B23", "Net Profit (MM)", "C23", "=IF(C11,C21-C7,-C4-C5)"),
    ]
    for lbl_cell, lbl, val_cell, val in rows:
        ws[lbl_cell] = lbl
        if val_cell:
            ws[val_cell] = val
    for a in ("C5", "C6", "C10", "C15", "C19"):
        ws[a].fill = GREEN
    for f in ("C21", "C23"):
        ws[f].fill = BLUE
        ws[f].font = BOLD
    ws["B23"].font = BOLD

    setup_sheet(
        wb,
        [
            "Vision Research (ClearView) — how to run this in the Crystal Ball Web add-in:",
            "",
            "Assumptions (green cells) — select the cell, click “+ Assumption”, then set:",
            "   C5  Testing Costs        ->  Uniform, Minimum 3, Maximum 5",
            "   C6  Marketing Costs      ->  Triangular, Min 12, Likeliest 16, Max 18",
            "   C10 Patients Cured       ->  Binomial, Trials 100, Probability 0.25",
            "   C15 Growth Rate          ->  Custom (Outcomes box, one per line):",
            "            0, 0.05, 0.75      (range 0%..5%, weight 75%)",
            "            -0.15, -0.05, 0.25 (range -15%..-5%, weight 25%)",
            "   C19 Market Penetration   ->  Normal, Mean 0.08, Std. Dev. 0.02",
            "",
            "Forecasts (blue cells):",
            "   C21 Gross Profit if Approved,  C23 Net Profit   ->  click “+ Forecast” on each",
            "",
            "Run settings (Run tab):  Trials 5000, Monte Carlo, fixed seed 999",
            "",
            "Expected result for Net Profit (matches Oracle Crystal Ball):",
            "   Certainty of Net Profit ≥ $0 ≈ 79%; certainty of ≥ $2MM ≈ 72%.",
            "",
            "Note: the add-in's Normal truncation (min 5%) matches Crystal Ball's tutorial, but",
            "truncation isn't exposed in the UI yet; results match closely without it too.",
        ],
    )
    wb.save(os.path.join(OUT, "vision_research.xlsx"))
    print("wrote examples/vision_research.xlsx")


def main():
    os.makedirs(OUT, exist_ok=True)
    futura()
    vision_research()


if __name__ == "__main__":
    main()
